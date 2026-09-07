from pathlib import Path
import unittest
from unittest.mock import Mock

from openpyxl import Workbook

from core.models import ProcessingRecord
from reporting.excel_report import (
    ReportWriter,
    dominant_collected_document_date,
    release_filenames_path_for_run,
    report_path_for_run,
)


class ReportingTests(unittest.TestCase):
    def test_report_filename_mirrors_run_folder(self):
        run_dir = Path("output") / "MIAP00_08-14-2026_22-10-03-027"
        self.assertEqual(
            report_path_for_run(run_dir).name,
            "Report_MIAP00_08-14-2026_22-10-03-027.xlsx",
        )

    def test_release_workbook_uses_majority_collected_decision_date(self):
        records = [
            ProcessingRecord(
                status="collected",
                docket=str(index),
                title="",
                release_date="",
                source_filename=f"source-{index}.pdf",
                source_url="",
                target_filename=f"LDC_SMD_{index}_{document_date}.pdf",
                document_date=document_date,
            )
            for index, document_date in enumerate(
                ["09032026", "09042026", "09042026", "09022026"],
                1,
            )
        ]
        run_dir = Path("output") / "MIAP00_09-07-2026_19-04-29-983"

        self.assertEqual(
            dominant_collected_document_date(records),
            "09042026",
        )
        self.assertEqual(
            release_filenames_path_for_run(run_dir, records).name,
            "MIAP00 09042026 Release Date Filenames.xlsx",
        )

    def test_release_workbook_date_tie_prefers_latest_decision_date(self):
        records = [
            ProcessingRecord(
                status="collected",
                docket="1",
                title="",
                release_date="",
                source_filename="source-1.pdf",
                source_url="",
                target_filename="LDC_SMD_1_09032026.pdf",
                document_date="09032026",
            ),
            ProcessingRecord(
                status="collected",
                docket="2",
                title="",
                release_date="",
                source_filename="source-2.pdf",
                source_url="",
                target_filename="LDC_SMD_2_09042026.pdf",
                document_date="09042026",
            ),
        ]

        self.assertEqual(
            dominant_collected_document_date(records),
            "09042026",
        )

    def test_release_workbook_is_not_named_without_collected_orders(self):
        duplicate = ProcessingRecord(
            status="duplicate",
            docket="1",
            title="",
            release_date="",
            source_filename="source.pdf",
            source_url="",
            target_filename="LDC_SMD_1_09042026.pdf",
            document_date="09042026",
        )
        self.assertIsNone(
            release_filenames_path_for_run(Path("output"), [duplicate])
        )

    def test_filenames_sheet_contains_only_final_collected_names(self):
        workbook = Workbook()
        workbook.remove(workbook.active)
        records = [
            ProcessingRecord(
                status="collected",
                docket="1",
                title="",
                release_date="",
                source_filename="source-1.pdf",
                source_url="",
                target_filename="LDC_SMD_1_08142026.pdf",
                counsel_references=[
                    "1: LDC_SMD_1counsel.html",
                    "10: 123456789",
                ],
            ),
            ProcessingRecord(
                status="duplicate",
                docket="2",
                title="",
                release_date="",
                source_filename="source-2.pdf",
                source_url="",
                target_filename="LDC_SMD_2_08142026.pdf",
            ),
            ProcessingRecord(
                status="collected",
                docket="3",
                title="",
                release_date="",
                source_filename="source-3.pdf",
                source_url="",
                target_filename="LDC_SMD_3_08142026.pdf",
            ),
            ProcessingRecord(
                status="content_duplicate",
                docket="4",
                title="",
                release_date="",
                source_filename="source-4.pdf",
                source_url="",
                target_filename="LDC_SMD_4_08142026.pdf",
            ),
        ]
        ReportWriter(Mock(), Mock())._add_filenames_sheet(workbook, records)
        sheet = workbook["Filenames"]
        self.assertEqual(
            [cell.value for cell in sheet["A"]],
            [
                "Main Document Filename",
                "LDC_SMD_1_08142026.pdf",
                "LDC_SMD_3_08142026.pdf",
            ],
        )
        self.assertEqual(
            [cell.value for cell in sheet["B"]],
            [
                "Counsel Filename / Recycled LNI",
                "1: LDC_SMD_1counsel.html\n10: 123456789",
                "",
            ],
        )
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.auto_filter.ref, "A1:B3")

    def test_release_workbook_uses_orders_sheet_and_filenames_mapping(self):
        from tempfile import TemporaryDirectory

        records = [
            ProcessingRecord(
                status="collected",
                docket="376000",
                title="",
                release_date="",
                source_filename="376000_20_01.pdf",
                source_url="",
                target_filename="LDC_SMD_376000_09042026.pdf",
                document_date="09042026",
                counsel_references=[
                    "166699: LDC_SMD_166699counsel.html",
                    "376000: LDC_SMD_376000counsel.html",
                ],
            )
        ]
        with TemporaryDirectory() as directory:
            writer = ReportWriter(Path(directory), Mock())
            path = writer._write_release_filenames_workbook(records)
            workbook = __import__("openpyxl").load_workbook(path)

        self.assertEqual(workbook.sheetnames, ["ORDERS"])
        sheet = workbook["ORDERS"]
        self.assertEqual(
            list(sheet.values),
            [
                ("Main Document Filename", "Counsel Filename / Recycled LNI"),
                (
                    "LDC_SMD_376000_09042026.pdf",
                    "166699: LDC_SMD_166699counsel.html\n"
                    "376000: LDC_SMD_376000counsel.html",
                ),
            ],
        )
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.auto_filter.ref, "A1:B2")


if __name__ == "__main__":
    unittest.main()
